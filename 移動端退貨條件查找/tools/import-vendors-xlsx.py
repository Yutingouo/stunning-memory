import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def clean_cell(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def extract_months(text):
    normalized = str(text or "").replace("個月以上", "個月").replace(" 個月", "個月")
    if "半年" in normalized:
        return 6
    if re.search(r"1\s*年|一年", normalized):
        return 12

    values = [int(item) for item in re.findall(r"(\d+)\s*個?月", normalized)]
    values.extend(int(item) for item in re.findall(r"[水粉液體膠囊粉錠肉鬆]\s*(\d+)", normalized))
    return max(values) if values else None


def infer_policy(condition, method):
    text = f"{condition or ''} {method or ''}".strip()
    if not text:
        return "unknown", None
    if any(keyword in text for keyword in ["禁退", "不可退"]):
        return "no_return", None

    months = extract_months(condition or "")
    if any(keyword in text for keyword in ["通知", "洽", "看商品", "滯銷", "瑕疵", "換貨", "可退", "無效期", "效期長", "當月"]):
        return "conditional", months
    return ("months", months) if months is not None else ("conditional", None)


def build_vendors(workbook_path):
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    vendors = []

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=2, max_col=9, values_only=True):
            code = clean_cell(row[0] if len(row) > 0 else "")
            delivery_vendor = clean_cell(row[1] if len(row) > 1 else "")
            name = clean_cell(row[2] if len(row) > 2 else "")
            condition = clean_cell(row[3] if len(row) > 3 else "")
            method = clean_cell(row[4] if len(row) > 4 else "")
            phone = clean_cell(row[5] if len(row) > 5 else "")
            fax = clean_cell(row[6] if len(row) > 6 else "")
            contact = clean_cell(row[7] if len(row) > 7 else "")
            note = clean_cell(row[8] if len(row) > 8 else "")

            if not any([code, delivery_vendor, name, condition, method, phone, fax, contact, note]):
                continue
            if code == "廠編" or name == "廠商名稱":
                continue
            if not code and not name:
                continue

            policy_type, min_months = infer_policy(condition, method)
            note_parts = []
            for label, value in [
                ("退貨方式", method),
                ("電話", phone),
                ("傳真", fax),
                ("聯絡人", contact),
                ("備註", note),
            ]:
                if value:
                    note_parts.append(f"{label}: {value}")

            if code.isdigit():
                code = code.zfill(3)

            vendors.append({
                "vendor_code": code,
                "delivery_vendor": delivery_vendor,
                "vendor_name": name or "(未填廠商名稱)",
                "min_months": min_months,
                "policy_type": policy_type,
                "condition_text": condition or method or "未填退貨條件",
                "return_method": method,
                "phone": phone,
                "fax": fax,
                "contact": contact,
                "note": "；".join(note_parts),
                "source_sheet": sheet.title,
            })

    return vendors


def main():
    parser = argparse.ArgumentParser(description="Import vendor return rules from an .xlsx workbook.")
    parser.add_argument("--input", required=True)
    parser.add_argument("--json", required=True)
    parser.add_argument("--js", required=True)
    args = parser.parse_args()

    vendors = build_vendors(args.input)
    if not vendors:
        raise RuntimeError("Excel returned no vendor rows.")

    json_text = json.dumps(vendors, ensure_ascii=False, indent=2)
    Path(args.json).write_text(json_text, encoding="utf-8")
    Path(args.js).write_text("window.VENDOR_DATA = " + json_text + ";\n", encoding="utf-8")
    print(f"已輸出 {len(vendors)} 筆廠商資料")


if __name__ == "__main__":
    main()
